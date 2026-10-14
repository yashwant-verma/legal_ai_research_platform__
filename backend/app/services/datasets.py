from __future__ import annotations
import csv, io, json, uuid
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from app.db.mongo import db
from app.evaluation.language_checks import language_issues
from app.schemas.dataset import DatasetCreate, DatasetQuestionIn, DatasetValidation
from app.schemas.factual import DOMAINS

def _now():
    return datetime.now(timezone.utc)

def validate_questions(questions: list[DatasetQuestionIn]) -> DatasetValidation:
    errors = []
    warnings = []
    ids = [q.question_id for q in questions]
    dup = [x for x, c in Counter(ids).items() if c > 1]
    if dup:
        errors.append(f"Duplicate question_id values: {', '.join(dup[:10])}")
    if not questions:
        errors.append('Dataset contains no questions.')
    for q in questions:
        if q.factual:
            f = q.factual
            if f.review_status != 'verified': warnings.append(f'{q.question_id}: factual labels are drafts; group review required.')
            if q.domain not in DOMAINS: errors.append(f'{q.question_id}: unsupported factual domain.')
            if not q.group_id: errors.append(f'{q.question_id}: factual rows require group_id.')
            if not q.answerable_from_kb: errors.append(f'{q.question_id}: factual v1 only supports answerable questions.')
            if not q.gold_chunk_ids or set(q.gold_chunk_hashes) != set(q.gold_chunk_ids):
                errors.append(f'{q.question_id}: factual rows need gold chunk IDs and text hashes.')
            if set(x.chunk_id for x in f.supporting_quotes) != set(q.gold_chunk_ids):
                errors.append(f'{q.question_id}: each gold chunk needs a supporting quote.')
        if q.answerable_from_kb and (not q.gold_document_ids) and (not q.gold_chunk_ids):
            warnings.append(f'{q.question_id}: no gold source labels; Hit@K cannot be scored reliably.')
        if q.gold_chunk_ids and set(q.gold_chunk_hashes)!=set(q.gold_chunk_ids):
            warnings.append(f'{q.question_id}: missing source text fingerprints; regenerate reviewed labels after reindexing.')
        if not q.reference_answer:
            warnings.append(f'{q.question_id}: reference_answer missing; answer correctness requires manual/external evaluation.')
    by_group={}
    factual_splits={}
    for q in questions:
        if q.factual:
            for cid in q.gold_chunk_ids:
                factual_splits.setdefault(cid,set()).add(q.factual.split)
        (errors if q.factual else warnings).extend(f'{q.question_id}: {issue}' for issue in language_issues(q.question,q.language))
        if q.group_id: by_group.setdefault(q.group_id,[]).append(q)
        if q.answerable_from_kb is False and (q.gold_document_ids or q.gold_chunk_ids):
            errors.append(f'{q.question_id}: unanswerable rows must not have positive gold evidence IDs.')
    for group,members in by_group.items():
        factual = [q for q in members if q.factual]
        if factual:
            if len(factual) != len(members): errors.append(f'{group}: cannot mix factual and legacy rows.')
            if {q.language for q in members} != {'en','hi','hinglish'}:
                errors.append(f'{group}: factual groups require exactly en, hi and hinglish.')
            common = {json.dumps(q.factual.model_dump(exclude={'accepted_aliases','reviewer','reviewed_at'}), sort_keys=True) for q in factual}
            if len(common) != 1: errors.append(f'{group}: translations must share factual labels, scope, review status and split.')
        if len({q.language for q in members})!=len(members): errors.append(f'{group}: duplicate language in question group.')
        if len({(tuple(q.gold_document_ids),tuple(q.gold_chunk_ids),q.answerable_from_kb) for q in members})>1:
            errors.append(f'{group}: translations must share gold IDs and answerability.')
    for cid, splits in factual_splits.items():
        if len(splits)>1: errors.append(f'{cid}: source chunk appears in both development and test splits.')
    langs = Counter((q.language for q in questions))
    groups = {q.group_id for q in questions if q.group_id}
    if {'en', 'hi', 'hinglish'}.issubset(langs) and groups:
        for g in sorted(groups):
            gl = {q.language for q in questions if q.group_id == g}
            if not {'en', 'hi', 'hinglish'}.issubset(gl):
                warnings.append(f'{g}: multilingual triplet incomplete.')
    return DatasetValidation(valid=not errors, errors=errors, warnings=warnings, capabilities={'retrieval_hit_at_k': all((not q.answerable_from_kb or bool(q.gold_document_ids or q.gold_chunk_ids) for q in questions)), 'reference_answer_available': all((bool(q.reference_answer) for q in questions)), 'factual_scoring_ready': bool(questions) and all(q.factual and q.factual.review_status=='verified' for q in questions), 'multilingual_analysis': len(langs) > 1})

def create_dataset(payload: DatasetCreate) -> dict:
    validation = validate_questions(payload.questions)
    if not validation.valid:
        raise ValueError('; '.join(validation.errors))
    if payload.benchmark and validation.warnings:
        raise ValueError('Benchmark requires review: '+'; '.join(validation.warnings))
    dataset_id = f'DS_{uuid.uuid4().hex[:12]}'
    doc = {'dataset_id': dataset_id, 'name': payload.name, 'version': payload.version, 'description': payload.description, 'benchmark': payload.benchmark, 'created_at': _now(), 'updated_at': _now(), 'validation': validation.model_dump()}
    db.datasets.insert_one(doc)
    if payload.questions:
        db.dataset_questions.insert_many([{'dataset_id': dataset_id, **q.model_dump()} for q in payload.questions])
    return get_dataset(dataset_id)

def get_dataset(dataset_id: str) -> dict:
    d = db.datasets.find_one({'dataset_id': dataset_id}, {'_id': 0})
    if not d:
        raise KeyError(dataset_id)
    qs = list(db.dataset_questions.find({'dataset_id': dataset_id}, {'_id': 0}))
    lc = Counter((q['language'] for q in qs))
    dc = Counter((q['domain'] for q in qs))
    return {**d, 'question_count': len(qs), 'language_counts': dict(lc), 'domain_counts': dict(dc)}

def list_datasets() -> list[dict]:
    return [get_dataset(x['dataset_id']) for x in db.datasets.find({}, {'dataset_id': 1, '_id': 0}).sort('created_at', -1)]

def get_questions(dataset_id: str) -> list[dict]:
    return list(db.dataset_questions.find({'dataset_id': dataset_id}, {'_id': 0}).sort('question_id', 1))

def delete_dataset(dataset_id: str) -> None:
    if db.experiments.count_documents({'dataset_id': dataset_id}) > 0:
        raise ValueError('Dataset has experiments and cannot be deleted.')
    db.dataset_questions.delete_many({'dataset_id': dataset_id})
    db.datasets.delete_one({'dataset_id': dataset_id})

def _parse_list(v: Any) -> list[str]:
    if v is None:
        return []
    if isinstance(v, list):
        return [str(x).strip() for x in v if str(x).strip()]
    s = str(v).strip()
    if not s:
        return []
    if s.startswith('['):
        try:
            return [str(x).strip() for x in json.loads(s)]
        except Exception:
            pass
    return [x.strip() for x in s.replace('|', ';').split(';') if x.strip()]

def rows_to_questions(rows: list[dict[str, Any]], mapping: dict[str, str] | None=None) -> list[DatasetQuestionIn]:
    mapping = mapping or {}
    out = []
    for i, row in enumerate(rows, 1):
        normalized = {}
        for k, v in row.items():
            normalized[mapping.get(k, k)] = v
        if isinstance(normalized.get('factual'), str):
            normalized['factual'] = json.loads(normalized['factual']) if normalized['factual'].strip() else None
        if isinstance(normalized.get('gold_authorities'), str):
            normalized['gold_authorities'] = json.loads(normalized['gold_authorities'] or '[]')
        if isinstance(normalized.get('gold_chunk_hashes'),str):
            normalized['gold_chunk_hashes']=json.loads(normalized['gold_chunk_hashes'] or '{}')
        normalized.setdefault('question_id', f'Q_{i:04d}')
        normalized.setdefault('domain', 'general_legal_procedure')
        normalized.setdefault('language', 'other')
        normalized['gold_document_ids'] = _parse_list(normalized.get('gold_document_ids') or normalized.get('gold_document_id'))
        normalized['gold_chunk_ids'] = _parse_list(normalized.get('gold_chunk_ids') or normalized.get('gold_chunk_id'))
        for b in ('answerable_from_kb', 'requires_multiple_sources'):
            if b in normalized and isinstance(normalized[b], str):
                normalized[b] = normalized[b].strip().lower() in {'1', 'true', 'yes', 'y'}
        out.append(DatasetQuestionIn.model_validate(normalized))
    return out

def parse_dataset_bytes(filename: str, data: bytes) -> list[dict[str, Any]]:
    ext = Path(filename).suffix.lower()
    if ext == '.csv':
        text = data.decode('utf-8-sig')
        return list(csv.DictReader(io.StringIO(text)))
    if ext in {'.json', '.jsonl'}:
        text = data.decode('utf-8')
        if ext == '.jsonl':
            return [json.loads(x) for x in text.splitlines() if x.strip()]
        obj = json.loads(text)
        return obj if isinstance(obj, list) else obj.get('questions', [])
    if ext == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        headers = [str(x or '').strip() for x in rows[0]]
        return [dict(zip(headers, r)) for r in rows[1:] if any((x is not None for x in r))]
    raise ValueError('Supported dataset formats: CSV, JSON, JSONL, XLSX')
